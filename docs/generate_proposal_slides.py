"""Create a seven-slide project proposal deck from the reviewed repository."""

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "assets"
OUTPUT = ROOT / "Computer_Vision_Autonomous_Patrol_Proposal_Presentation.pptx"

NAVY = RGBColor(10, 19, 26)
PANEL = RGBColor(18, 40, 51)
TEAL = RGBColor(96, 215, 216)
WHITE = RGBColor(235, 245, 246)
MUTED = RGBColor(163, 190, 198)
GOLD = RGBColor(255, 209, 102)
RED = RGBColor(244, 114, 124)


def rectangle(slide, x, y, w, h, color, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background() if line is None else None
    return shape


def text(slide, value, x, y, w, h, size=20, color=WHITE, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        if align is not None:
            p.alignment = align
        p.space_after = Pt(4)
        for run in p.runs:
            run.font.name = "Aptos"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color
    return box


def base_slide(prs, title, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rectangle(slide, 0, 0, 13.333, 7.5, NAVY)
    rectangle(slide, 0, 0, 0.12, 7.5, TEAL)
    text(slide, title, 0.48, 0.24, 12.2, 0.55, 27, WHITE, True)
    rectangle(slide, 0.5, 0.87, 12.35, 0.025, TEAL)
    text(slide, "COMPUTER VISION AUTONOMOUS PATROL SIMULATOR", 0.52, 7.13, 10.8, 0.2, 9, MUTED)
    text(slide, str(number), 12.32, 7.08, 0.4, 0.25, 10, TEAL, True, PP_ALIGN.RIGHT)
    return slide


def panel(slide, x, y, w, h, heading, body, accent=TEAL):
    rectangle(slide, x, y, w, h, PANEL)
    rectangle(slide, x, y, 0.055, h, accent)
    text(slide, heading, x + 0.2, y + 0.12, w - 0.4, 0.38, 18, accent, True)
    text(slide, body, x + 0.2, y + 0.55, w - 0.4, h - 0.68, 15, WHITE)


def picture_fit(slide, path, x, y, w, h):
    with Image.open(path) as img:
        aspect = img.width / img.height
    box_aspect = w / h
    if aspect > box_aspect:
        draw_w, draw_h = w, w / aspect
    else:
        draw_w, draw_h = h * aspect, h
    slide.shapes.add_picture(str(path), Inches(x + (w - draw_w) / 2),
                             Inches(y + (h - draw_h) / 2),
                             width=Inches(draw_w), height=Inches(draw_h))


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Computer Vision Autonomous Patrol Simulator Project Proposal"
    prs.core_properties.subject = "Academic project concept and validation proposal"

    slide = base_slide(prs, "Project proposal", 1)
    text(slide, "Computer Vision\nAutonomous Patrol Simulator", 0.75, 1.55, 7.4, 2.2, 36, WHITE, True)
    text(slide, "From a camera frame to an explainable\nsimulated navigation decision", 0.79, 4.0, 7.4, 1.0, 22, TEAL)
    rectangle(slide, 8.75, 1.6, 3.75, 4.45, PANEL)
    text(slide, "PROTOTYPE STATUS", 9.02, 1.98, 3.2, 0.4, 18, GOLD, True)
    text(slide, "Working GUI and navigation pipeline\n\n6 headless tests passed\n\nReal-footage evaluation proposed", 9.02, 2.68, 3.1, 2.45, 18, WHITE)
    text(slide, "Software-only educational system  •  3 October 2026", 0.79, 6.47, 9.5, 0.4, 15, MUTED)

    slide = base_slide(prs, "Problem and project question", 2)
    panel(slide, 0.65, 1.28, 5.75, 3.6, "The problem",
          "Object detection identifies what appears in an image. It does not, by itself, show which route is less risky or explain a navigation command.")
    panel(slide, 6.75, 1.28, 5.9, 3.6, "Our question",
          "How can monocular detections, image geometry, and transparent risk rules produce stable simulated decisions?", GOLD)
    text(slide, "Goal: make every step from object box to steering recommendation visible and testable.",
         0.8, 5.45, 11.7, 0.85, 23, WHITE, True)

    slide = base_slide(prs, "Project objectives and scope", 3)
    items = [
        ("01", "Detect and track", "Selected people, vehicles, bicycles, buses, trucks, and chairs."),
        ("02", "Understand scene layout", "Zones, projected corridor, approximate proximity, and apparent approach."),
        ("03", "Recommend an action", "Compare left and right risk; choose straight, turn, or stop."),
        ("04", "Explain and evaluate", "Show the reason in the GUI; replay scenes and inspect failures."),
    ]
    for i, (number, heading, body) in enumerate(items):
        y = 1.22 + i * 1.25
        rectangle(slide, 0.69, y, 11.95, 1.06, PANEL)
        text(slide, number, 0.93, y + 0.13, 0.68, 0.62, 27, TEAL, True)
        text(slide, heading, 1.75, y + 0.08, 3.15, 0.4, 19, WHITE, True)
        text(slide, body, 4.92, y + 0.13, 7.25, 0.7, 16, MUTED)
    text(slide, "Scope: desktop simulation only; no motors, measured depth, or safety claim.",
         0.83, 6.5, 11.8, 0.35, 16, GOLD)

    slide = base_slide(prs, "How the system works", 4)
    rectangle(slide, 0.68, 1.22, 11.98, 3.67, RGBColor(255, 255, 255))
    picture_fit(slide, ASSETS / "system_flow.png", 0.81, 1.35, 11.7, 3.34)
    text(slide, "Frame → objects → tracking → geometry → risk → route choice → steering and wheel simulation",
         0.88, 5.17, 11.5, 0.84, 21, WHITE, True)
    text(slide, "The dashboard displays the frame, estimated proximity, corridor risks, selected action, and reason.",
         0.88, 6.17, 11.5, 0.45, 16, MUTED)

    slide = base_slide(prs, "Computer vision contribution", 5)
    panel(slide, 0.66, 1.2, 5.82, 2.07, "Perception",
          "YOLO11n returns classes, confidence, and bounding boxes. IoU matching assigns track IDs and box-area trends suggest approach.")
    panel(slide, 6.8, 1.2, 5.84, 2.07, "Spatial reasoning",
          "Box centers select image zones. Box height gives an approximate distance; a center corridor and lower-frame position add risk.", GOLD)
    panel(slide, 0.66, 3.55, 5.82, 2.07, "Scene risk",
          "Risk combines proximity, size, class, confidence, location, and approach. Zone risks are combined for each side.")
    panel(slide, 6.8, 3.55, 5.84, 2.07, "Explainability",
          "The GUI shows the primary threat, its estimated distance and risk, left/right corridor totals, and the selected action.", GOLD)
    text(slide, "Distance and risk are heuristics, not calibrated physical measurements.",
         0.83, 6.22, 11.6, 0.42, 17, RED, True)

    slide = base_slide(prs, "Current prototype demonstration", 6)
    picture_fit(slide, ASSETS / "dashboard_turn_left.png", 0.64, 1.14, 7.1, 5.75)
    text(slide, "CONTROLLED TURN SCENE", 7.95, 1.45, 4.45, 0.45, 19, TEAL, True)
    text(slide, "A center car and a right-side person raise risk on the right. The simulator selects TURN LEFT, with the left wheel slower than the right.",
         7.95, 2.1, 4.38, 2.3, 21, WHITE)
    rectangle(slide, 7.94, 4.9, 4.47, 1.35, PANEL)
    text(slide, "Figure provenance", 8.12, 5.02, 4.0, 0.35, 15, GOLD, True)
    text(slide, "Real GUI capture with injected test boxes over synthetic video; not a YOLO accuracy result.",
         8.12, 5.4, 3.97, 0.68, 13, MUTED)

    slide = base_slide(prs, "Validation and next steps", 7)
    panel(slide, 0.67, 1.27, 5.77, 3.94, "Evidence already available",
          "Working webcam/video dashboard.\nSix navigation tests passed.\nControlled clear, turn, and stop screenshots.\nYOLO model loaded on CPU.")
    panel(slide, 6.79, 1.27, 5.84, 3.94, "Proposed evaluation",
          "Use real footage with configured classes.\nRecord detection and tracking behavior.\nMeasure FPS and inference time.\nResolve pause, stop-delay, and corridor-display inconsistencies.", GOLD)
    text(slide, "Outcome: a reproducible teaching and research demonstration with clearly reported limits.",
         0.84, 5.65, 11.6, 0.82, 22, WHITE, True)

    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
