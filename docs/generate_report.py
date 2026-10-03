"""Render the Markdown project documents as styled Word files with figures."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PROJECT_REPORT.md"
OUTPUT = ROOT / "Computer_Vision_Object_Detection_Patrol_System.docx"
TEAL = RGBColor(20, 89, 106)
INK = RGBColor(32, 49, 59)


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    tc_pr.append(element)


def add_inline(paragraph, value: str) -> None:
    """Add the small Markdown subset used by the source report."""
    value = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", value)
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", value):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(8.5)
        else:
            paragraph.add_run(part)


def build_document(source: Path, output: Path, *, cover_title: str, subtitle_text: str,
                   header_text: str, strapline: str, metadata_line: str,
                   document_subject: str) -> Path:
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.66)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.33)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.13
    for style_name, size in [("Title", 26), ("Heading 1", 15), ("Heading 2", 11.5)]:
        style = doc.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = TEAL
        style.paragraph_format.space_before = Pt(13 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.text = header_text
    header.runs[0].font.size = Pt(7)
    header.runs[0].font.color.rgb = TEAL
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Computer Vision Autonomous Patrol Simulator  ·  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(7)
        run.font.color.rgb = TEAL

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(cover_title)
    subtitle = doc.add_paragraph(subtitle_text)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(15)
    subtitle.runs[0].font.color.rgb = TEAL
    doc.add_paragraph(strapline).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(metadata_line).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(str(ROOT / "docs" / "assets" / "system_flow.png"), width=Inches(6.8))
    doc.add_page_break()

    doc.add_heading("Contents", 1)
    for line in lines:
        if line.startswith("## "):
            para = doc.add_paragraph(line[3:])
            para.paragraph_format.space_after = Pt(3)
    doc.add_page_break()

    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        if line.startswith("# ") or line.startswith("**") and line.split(":", 1)[0] in (
                "**Engineering analysis and demonstration report**  ", "**Prepared for", "**Project area",
                "**Status", "**Date", "**Repository", "**Reviewed", "**Application type"):
            index += 1
            continue
        if line.startswith("```"):
            if in_code:
                para = doc.add_paragraph(style="No Spacing")
                para.paragraph_format.space_before = Pt(4)
                para.paragraph_format.space_after = Pt(8)
                for n, code_line in enumerate(code_lines):
                    run = para.add_run(code_line + ("\n" if n < len(code_lines) - 1 else ""))
                    run.font.name = "Consolas"
                    run.font.size = Pt(7.5)
                code_lines.clear()
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 2)
        elif line.startswith("!["):
            match = re.match(r"!\[[^]]*\]\(([^)]+)\)", line)
            if match:
                image_path = source.parent / match.group(1)
                if not image_path.exists():
                    raise FileNotFoundError(image_path)
                para = doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.keep_with_next = True
                width = 5.0 if image_path.name == "frame_pipeline.png" else 6.85
                para.add_run().add_picture(str(image_path), width=Inches(width))
        elif line.startswith("| ") and index + 1 < len(lines) and lines[index + 1].startswith("|---"):
            rows = []
            while index < len(lines) and lines[index].startswith("| "):
                if not lines[index].startswith("|---"):
                    rows.append([v.strip() for v in lines[index].strip().strip("|").split("|")])
                index += 1
            if rows:
                table = doc.add_table(rows=0, cols=len(rows[0]))
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                for row_number, values in enumerate(rows):
                    cells = table.add_row().cells
                    for col_number, value in enumerate(values):
                        cells[col_number].text = ""
                        add_inline(cells[col_number].paragraphs[0], value)
                        if row_number == 0:
                            shade(cells[col_number], "DCECF0")
                            for run in cells[col_number].paragraphs[0].runs:
                                run.bold = True
                        elif row_number % 2 == 0:
                            shade(cells[col_number], "F4F8F9")
                doc.add_paragraph()
            continue
        elif re.match(r"\d+\. ", line):
            para = doc.add_paragraph(style="List Number")
            add_inline(para, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            para = doc.add_paragraph(style="List Bullet")
            add_inline(para, line[2:])
        elif line.startswith("**Figure "):
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(10)
            add_inline(para, line)
            for run in para.runs:
                run.font.size = Pt(8)
                run.font.color.rgb = TEAL
        elif line.strip():
            para = doc.add_paragraph()
            add_inline(para, line.rstrip("  "))
        index += 1

    doc.core_properties.title = cover_title.replace("\n", " ") + " — " + subtitle_text
    doc.core_properties.subject = document_subject
    doc.core_properties.keywords = "computer vision, YOLO, object detection, navigation, PySide6"
    doc.save(output)
    return output


def build_report() -> Path:
    return build_document(
        SOURCE, OUTPUT,
        cover_title="Computer Vision\nAutonomous Patrol Simulator",
        subtitle_text="Engineering analysis and demonstration report",
        header_text="CV PROJECT  /  ENGINEERING REPORT",
        strapline="Code review · Computer vision pipeline · GUI guide · Reproducible outcomes",
        metadata_line="Reviewed 3 October 2026  |  Software-only educational prototype",
        document_subject="Code analysis, CV pipeline, GUI, screenshots, execution and validation",
    )


if __name__ == "__main__":
    print(build_report())
